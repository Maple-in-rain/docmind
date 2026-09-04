"""Word 解析：python-docx 提取段落 + 表格。

表格的每一行用 " | " 连接单元格，保留表格信息的语义。
"""

from docx import Document

from .base import BaseParser, ParsedDoc


class DocxParser(BaseParser):
    def parse(self, file_path) -> ParsedDoc:
        doc = Document(str(file_path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return ParsedDoc(text="\n".join(parts), meta={"format": "docx"})
