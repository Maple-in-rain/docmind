"""解析器单测：4 种格式 + GBK 编码坑 + 注册表分发。

PDF / Word 测试文件在测试内程序化生成，避免提交二进制 fixture。
"""

import pytest

from app.parsers.docx_parser import DocxParser
from app.parsers.md_parser import MarkdownParser
from app.parsers.pdf_parser import PdfParser
from app.parsers.registry import get_parser
from app.parsers.txt_parser import TxtParser


# ---- TXT ----

def test_txt_utf8(tmp_path):
    p = tmp_path / "utf8.txt"
    p.write_text("你好，世界。\n第二行。", encoding="utf-8")
    parsed = TxtParser().parse(p)
    assert "你好，世界。" in parsed.text
    assert "第二行。" in parsed.text


def test_txt_gbk(tmp_path):
    """GBK 编码是中文 Windows 常见坑，必须自动探测而非硬编码 UTF-8"""
    p = tmp_path / "gbk.txt"
    p.write_bytes("中文内容测试".encode("gbk"))
    parsed = TxtParser().parse(p)
    assert parsed.text == "中文内容测试"
    assert parsed.meta["encoding"].lower() in ("gbk", "gb18030", "gb2312")


# ---- Markdown ----

def test_md_strips_format_and_keeps_code(tmp_path):
    p = tmp_path / "doc.md"
    p.write_text(
        "# 标题一\n\n正文**加粗**内容\n\n## 标题二\n\n- 列表项\n\n```python\nprint('代码保留')\n```\n",
        encoding="utf-8",
    )
    parsed = MarkdownParser().parse(p)
    assert parsed.headings == [(1, "标题一"), (2, "标题二")]
    assert "**" not in parsed.text  # 格式符号已剥离
    assert "print('代码保留')" in parsed.text  # 代码块原样保留


# ---- Word ----

def test_docx_paragraphs_and_tables(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_paragraph("第一段内容")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "单元格A"
    table.cell(0, 1).text = "单元格B"
    p = tmp_path / "t.docx"
    doc.save(str(p))

    parsed = DocxParser().parse(p)
    assert "第一段内容" in parsed.text
    assert "单元格A | 单元格B" in parsed.text


# ---- PDF ----

def test_pdf_text_extraction(tmp_path):
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Hello PDF world")
    p = tmp_path / "t.pdf"
    doc.save(str(p))
    doc.close()

    parsed = PdfParser().parse(p)
    assert "Hello PDF world" in parsed.text


# ---- 注册表 ----

def test_registry_dispatches_by_extension(tmp_path):
    md = tmp_path / "a.md"
    md.write_text("# hi", encoding="utf-8")
    assert isinstance(get_parser(md), MarkdownParser)

    txt = tmp_path / "b.TXT"  # 大写扩展名也要识别
    txt.write_text("x", encoding="utf-8")
    assert isinstance(get_parser(txt), TxtParser)


def test_registry_rejects_unsupported(tmp_path):
    p = tmp_path / "a.xyz"
    p.write_text("x", encoding="utf-8")
    with pytest.raises(ValueError, match="不支持的文件格式"):
        get_parser(p)
